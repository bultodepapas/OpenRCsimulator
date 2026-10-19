# Generate the current Word snapshot from docs/FUEL-SYSTEM.md.
# Authoring dependency python-docx==1.2.0; use the managed artifact runtime.
from pathlib import Path
import re, os, hashlib, json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

root = Path(__file__).resolve().parents[2]
source = root / 'docs/FUEL-SYSTEM.md'
output = root / 'docs/research/fuel-systems/FUEL-12/fuel-system-standard.docx'
output.parent.mkdir(parents=True,exist_ok=True)
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.75)
section.footer_distance = Inches(0.3)
styles = doc.styles
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Number']:
    style = styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.10
    style.paragraph_format.widow_control = True
    for child in list(style.element.xpath('./w:pPr/w:pBdr')):
        child.getparent().remove(child)
styles['Title'].font.size = Pt(23)
styles['Title'].font.bold = True
styles['Title'].paragraph_format.space_after = Pt(12)
for name, size, before, after in [('Heading 1', 15, 13, 7), ('Heading 2', 12.5, 10, 5)]:
    st = styles[name]
    st.font.size, st.font.bold = Pt(size), True
    st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(before), Pt(after)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.keep_together = True

def add_inline(p, text, small=False):
    tokens = re.split(r'(\[[^\]]+\]\([^\)]+\)|\*\*.+?\*\*|`[^`]+`)', text)
    for token in tokens:
        if not token:
            continue
        link = re.fullmatch(r'\[([^\]]+)\]\(([^\)]+)\)', token)
        if link:
            label, url = link.groups()
            if not url.startswith(('https://', 'http://')):
                path, sep, anchor = url.partition('#')
                target = (source.parent / path).resolve() if path else source
                url = os.path.relpath(target, output.parent).replace('\\', '/') + (sep+anchor if sep else '')
            relation = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
            el = OxmlElement('w:hyperlink'); el.set(qn('r:id'), relation)
            run = OxmlElement('w:r'); pr = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '25475B'); pr.append(color)
            underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); pr.append(underline)
            run.append(pr); tx = OxmlElement('w:t'); tx.text = label; run.append(tx); el.append(run); p._p.append(el)
        else:
            run = p.add_run(token[2:-2] if token.startswith('**') else token[1:-1] if token.startswith('`') else token)
            if token.startswith('**'): run.bold = True
            if token.startswith('`'):
                run.font.name = 'Consolas'
                run.font.size = Pt(10 if not small else 9.5)
            elif small: run.font.size = Pt(10)

def table(rows):
    count = len(rows[0])
    t = doc.add_table(rows=1, cols=count)
    t.alignment, t.autofit = WD_TABLE_ALIGNMENT.CENTER, False
    widths = [1.65, 1.18, 1.4, 2.77] if count == 4 else [1.22, 1.08, 4.70]
    for c,w in zip(t.columns,widths): c.width = Inches(w)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        el = OxmlElement('w:'+name); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
    pr.append(borders)
    for index, values in enumerate(rows):
        cells = t.rows[0].cells if index == 0 else t.add_row().cells
        rowpr = t.rows[index]._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); rowpr.append(cant)
        if index == 0:
            repeat = OxmlElement('w:tblHeader'); rowpr.append(repeat)
        for ci,(cell,value) in enumerate(zip(cells, values)):
            cell.width = Inches(widths[ci]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'DCE5EB' if index == 0 else 'F3F6F8' if index%2 else 'FFFFFF'); props.append(shade)
            margin = OxmlElement('w:tcMar')
            for side,number in [('top',80),('bottom',80),('left',95),('right',95)]:
                node = OxmlElement('w:'+side); node.set(qn('w:w'),str(number)); node.set(qn('w:type'),'dxa'); margin.append(node)
            props.append(margin)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0); p.paragraph_format.line_spacing = 1.06
            p.paragraph_format.keep_with_next = False
            if count == 3 and ci == 1: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, value, small=True)
            if index == 0:
                for r in p.runs: r.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

lines = source.read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1; continue
    if line.startswith('|'):
        rows=[]
        while i < len(lines) and lines[i].strip().startswith('|'):
            values=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',v) for v in values): rows.append(values)
            i += 1
        table(rows); continue
    heading = re.match(r'^(#{1,3}) (.+)$', line)
    if heading:
        level = len(heading.group(1))
        doc.add_paragraph(heading.group(2), 'Title' if level==1 else 'Heading '+str(level-1))
    elif re.match(r'^\d+\. ',line):
        p=doc.add_paragraph(style='List Number'); add_inline(p,re.sub(r'^\d+\. ','',line))
        p.paragraph_format.keep_with_next = i+1 < len(lines) and bool(re.match(r'^\d+\. ',lines[i+1].strip()))
    else:
        p=doc.add_paragraph(); add_inline(p,line)
        if i+1 < len(lines) and not lines[i+1].strip() and i+2 < len(lines) and re.match(r'^\d+\. ',lines[i+2].strip()):
            p.paragraph_format.keep_with_next = True
    i += 1

footer=section.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run('OpenRC fuel standard  |  '); run.font.name='Arial'; run.font.size=Pt(9)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
doc.core_properties.title='Fuel system engineering standard'
doc.core_properties.subject='Current fixed wing implementation and aircraft admission contract'
doc.core_properties.author='OpenRC simulator'
doc.core_properties.comments='Source docs/FUEL-SYSTEM.md SHA256 '+hashlib.sha256(source.read_bytes()).hexdigest()
doc.save(output)
print(json.dumps({'file':str(output),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'bytes':output.stat().st_size}))
